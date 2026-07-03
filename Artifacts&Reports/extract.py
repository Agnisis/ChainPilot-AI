import sys, subprocess
try: import docx
except ImportError: subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx'])
try: import pptx
except ImportError: subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-pptx'])

import docx, pptx
import os

folder = r'c:\Users\agnis\OneDrive\Desktop\My Workspace\SCMAi\Artifacts&Reports'
try:
    doc = docx.Document(os.path.join(folder, 'MSc_DS_Project_Report_Format.docx'))
    with open(os.path.join(folder, 'report_dump.txt'), 'w', encoding='utf-8') as f:
        for p in doc.paragraphs: f.write(p.text + '\n')
except Exception as e: print('Docx error:', e)

def extract_ppt(name, out):
    try:
        prs = pptx.Presentation(os.path.join(folder, name))
        with open(os.path.join(folder, out), 'w', encoding='utf-8') as f:
            for s in prs.slides:
                for sh in s.shapes:
                    if hasattr(sh, 'text'): f.write(sh.text + '\n')
    except Exception as e: print(f'PPTX error {name}:', e)

extract_ppt('SupplyChain_FirstReview_Recolored-1.pptx', 'ppt1_dump.txt')
extract_ppt('SupplyChain_Intelligence_Platform_PPT.pptx', 'ppt2_dump.txt')
extract_ppt('ThirdReview_ShilpiSen_Final (1).pptx', 'ppt3_dump.txt')
print('Extraction complete.')
